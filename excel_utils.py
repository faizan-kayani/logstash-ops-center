"""Import/export helpers for the Excel Sheet workspace (see excel_sheet table
comment in db.py). This has nothing to do with SSH/Logstash monitoring --
it's a small self-contained spreadsheet feature that lives entirely in this
dashboard's own database.

Grid shape used everywhere (DB storage, the editor's JSON API, and both
directions of .xlsx conversion below):

    {"rows": [[{"v": <str>, "bg": <"#rrggbb" or None>, "fg": <"#rrggbb" or None>}, ...], ...]}

"v" is the cell text, "bg"/"fg" are optional background/text colors.
"""

import io

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

MAX_IMPORT_ROWS = 2000
MAX_IMPORT_COLS = 200


def _clean_hex(value):
    """openpyxl hands back ARGB like 'FFRRGGBB' (or a non-hex theme/indexed
    reference for untouched cells); normalizes to a plain '#rrggbb' for the
    browser's <input type=color>, or None if there's nothing usable."""
    if not value or not isinstance(value, str):
        return None
    hexval = value.lstrip("#")
    if len(hexval) == 8:
        hexval = hexval[2:]
    if len(hexval) != 6:
        return None
    return f"#{hexval.lower()}"


def workbook_to_grid(file_stream):
    """Reads an uploaded .xlsx file's active sheet into our grid dict.
    Raises ValueError on anything that isn't a readable workbook."""
    try:
        wb = load_workbook(file_stream, data_only=True)
    except Exception as exc:
        raise ValueError(f"Could not read this file as an Excel workbook: {exc}")

    ws = wb.active
    max_row = min(ws.max_row or 1, MAX_IMPORT_ROWS)
    max_col = min(ws.max_column or 1, MAX_IMPORT_COLS)

    rows = []
    for r in range(1, max_row + 1):
        row = []
        for c in range(1, max_col + 1):
            cell = ws.cell(row=r, column=c)
            value = cell.value
            text = "" if value is None else str(value)
            bg = None
            if cell.fill is not None and cell.fill.fgColor is not None and cell.fill.patternType == "solid":
                bg = _clean_hex(getattr(cell.fill.fgColor, "rgb", None))
            fg = None
            if cell.font is not None and cell.font.color is not None:
                fg = _clean_hex(getattr(cell.font.color, "rgb", None))
            row.append({"v": text, "bg": bg, "fg": fg})
        rows.append(row)

    if not rows:
        rows = [[{"v": "", "bg": None, "fg": None} for _ in range(6)] for _ in range(10)]

    return {"rows": rows}


def grid_to_workbook_bytes(grid):
    """Serializes our grid dict back into a real .xlsx file (in-memory) for
    download -- includes cell colors, not just text."""
    wb = Workbook()
    ws = wb.active

    for r, row in enumerate(grid.get("rows", []), start=1):
        for c, cell in enumerate(row, start=1):
            target = ws.cell(row=r, column=c, value=cell.get("v") or "")
            bg = cell.get("bg")
            fg = cell.get("fg")
            if bg:
                hexval = bg.lstrip("#").upper()
                target.fill = PatternFill(start_color=hexval, end_color=hexval, fill_type="solid")
            if fg:
                hexval = fg.lstrip("#").upper()
                target.font = Font(color=hexval)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
